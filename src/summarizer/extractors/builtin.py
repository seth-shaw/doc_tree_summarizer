"""Built-in text, HTML, DOCX, and PDF extractors."""

import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
from bs4 import BeautifulSoup
from docx import Document

from .base import Extracted

TEXT = {'.txt', '.md', '.markdown', '.csv', '.json', '.xml', '.rst', '.log'}
HTML = {'.html', '.htm'}

def _ocr_pdf(path: Path, target: Path, ocr: dict, mode: str) -> None:
    """Run OCR through the selected backend, writing the result to *target*."""
    backend = ocr.get('backend', 'docker')
    options = ['--rotate-pages', '--deskew', '--jobs', '2', '-l', str(ocr.get('language', 'eng')),
               '--force-ocr' if mode == 'always' else '--skip-text']
    if backend == 'docker':
        if not shutil.which('docker'):
            raise RuntimeError(
                'PDF OCR needs Docker with Compose; install and start Docker Desktop'
            )
        default_compose = Path(__file__).resolve().parents[3] / 'compose.yaml'
        compose_file = Path(ocr.get('compose_file', default_compose)).expanduser().resolve()
        if not compose_file.is_file():
            raise RuntimeError(
                'OCR Compose file not found; set extraction.ocr.compose_file '
                'to its absolute path'
            )
        command = ['docker', 'compose', '-f', str(compose_file), 'run', '--rm', '--no-deps', '-T',
                   'ocrmypdf', *options, '-', '-']
        # Stream PDF bytes directly: no host mounts or container-owned output files.
        with path.open('rb') as source, target.open('wb') as output:
            result = subprocess.run(
                command, check=False, stdin=source, stdout=output,
                stderr=subprocess.PIPE, text=True, timeout=3600,
            )
    elif backend == 'native':
        if not shutil.which('ocrmypdf'):
            raise RuntimeError(
                'PDF OCR needs native ocrmypdf; install it or set '
                'extraction.ocr.backend to docker'
            )
        result = subprocess.run(
            ['ocrmypdf', *options, str(path), str(target)],
            check=False, capture_output=True, text=True, timeout=3600,
        )
    else:
        raise ValueError('extraction.ocr.backend must be docker or native')
    if result.returncode != 0:
        raise RuntimeError('OCR failed: ' + result.stderr[-1500:])

def extract(path: Path, config: dict) -> Extracted:
    """Extract supported document content using built-in libraries."""
    suffix = path.suffix.lower()
    if suffix in TEXT:
        return Extracted(path.read_text(encoding='utf-8', errors='replace'), 'builtin:text')
    if suffix in HTML:
        soup = BeautifulSoup(path.read_text(encoding='utf-8', errors='replace'), 'html.parser')
        for element in soup(['script', 'style', 'noscript']):
            element.decompose()
        return Extracted(soup.get_text('\n', strip=True), 'builtin:html')
    if suffix == '.docx':
        doc = Document(str(path))
        pieces = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                pieces.append(' | '.join(cell.text.replace('\n', ' ') for cell in row.cells))
        return Extracted(
            '\n\n'.join(pieces), 'builtin:docx',
            ['Embedded images and textboxes may not be extracted.'],
        )
    if suffix == '.pdf':
        return _extract_pdf(path, config.get('ocr', {}))
    raise ValueError(f'Unsupported builtin format: {suffix}')


def _extract_pdf(path: Path, ocr: dict) -> Extracted:
    """Extract PDF pages and optionally OCR pages with little native text."""
    mode = ocr.get('mode', 'auto')
    if mode not in ('off', 'auto', 'always'):
        raise ValueError('extraction.ocr.mode must be off, auto, or always')
    with fitz.open(path) as document:
        pages = [page.get_text('text').strip() for page in document]

    threshold = int(ocr.get('min_chars_per_page', 40))
    has_sparse_page = any(len(page) < threshold for page in pages)
    needs_ocr = mode == 'always' or (mode == 'auto' and has_sparse_page)
    warnings = []
    if needs_ocr:
        with tempfile.TemporaryDirectory(prefix='local-ai-ocr-') as temp_dir:
            target = Path(temp_dir) / 'ocr.pdf'
            _ocr_pdf(path, target, ocr, mode)
            with fitz.open(target) as document:
                pages = [page.get_text('text').strip() for page in document]
        warnings.append('OCR applied; verify recognition accuracy.')

    page_text = '\n\n'.join(
        f'[Page {index + 1}]\n{text}'
        for index, text in enumerate(pages)
        if text
    )
    return Extracted(page_text, 'builtin:pdf', warnings)
